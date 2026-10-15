"""
Resume text extraction service.
Supports PDF (PyMuPDF primary, pdfplumber fallback) and DOCX.
"""
import io
import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)


class ResumeParserError(Exception):
    """Raised when resume text cannot be extracted"""
    pass


class ResumeParser:
    """Extracts and cleans text from PDF and DOCX resume files"""

    MIN_CONTENT_LENGTH = 100  # Minimum characters expected from a real resume

    def parse(self, file_content: bytes, filename: str) -> str:
        """
        Parse resume from bytes based on file extension.
        
        Args:
            file_content: Raw bytes of the uploaded file
            filename: Original filename (used to determine file type)
            
        Returns:
            Cleaned extracted text
            
        Raises:
            ResumeParserError: If text cannot be extracted
        """
        ext = Path(filename).suffix.lower()
        
        if ext == ".pdf":
            text = self.parse_pdf(file_content)
        elif ext == ".docx":
            text = self.parse_docx(file_content)
        else:
            raise ResumeParserError(f"Unsupported file type: {ext}")
        
        cleaned = self._clean_text(text)
        
        if len(cleaned) < self.MIN_CONTENT_LENGTH:
            raise ResumeParserError(
                "Could not extract sufficient text from resume. "
                "If this is a scanned PDF, please use a text-based PDF or DOCX."
            )
        
        return cleaned

    def parse_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF using PyMuPDF with pdfplumber fallback"""
        # Primary: PyMuPDF (faster, better layout handling)
        try:
            return self._parse_pdf_pymupdf(file_content)
        except Exception as e:
            logger.warning(f"PyMuPDF failed, falling back to pdfplumber: {e}")
        
        # Fallback: pdfplumber
        try:
            return self._parse_pdf_pdfplumber(file_content)
        except Exception as e:
            logger.error(f"pdfplumber also failed: {e}")
            raise ResumeParserError(
                "Unable to extract text from PDF. "
                "The file may be corrupted or image-based."
            )

    def _parse_pdf_pymupdf(self, file_content: bytes) -> str:
        """Parse PDF using PyMuPDF"""
        import fitz  # PyMuPDF
        
        file_buffer = io.BytesIO(file_content)
        doc = fitz.open(stream=file_buffer, filetype="pdf")
        
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text)
        
        doc.close()
        
        full_text = "\n".join(text_parts)
        
        if len(full_text.strip()) < 50:
            raise ResumeParserError("Extracted text too short — may be image-based PDF")
        
        return full_text

    def _parse_pdf_pdfplumber(self, file_content: bytes) -> str:
        """Parse PDF using pdfplumber (fallback)"""
        import pdfplumber
        
        file_buffer = io.BytesIO(file_content)
        text_parts = []
        
        with pdfplumber.open(file_buffer) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        
        return "\n".join(text_parts)

    def parse_docx(self, file_content: bytes) -> str:
        """Extract text from DOCX using python-docx"""
        try:
            from docx import Document
            
            file_buffer = io.BytesIO(file_content)
            doc = Document(file_buffer)
            
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        text_parts.append(row_text)
            
            return "\n".join(text_parts)
        except Exception as e:
            logger.error(f"DOCX parsing failed: {e}")
            raise ResumeParserError(f"Unable to extract text from DOCX: {str(e)}")

    def _clean_text(self, text: str) -> str:
        """Clean and normalize extracted text"""
        if not text:
            return ""
        
        # Normalize whitespace
        # Replace multiple spaces with single space
        text = re.sub(r' {2,}', ' ', text)
        
        # Normalize line endings
        text = text.replace('\r\n', '\n').replace('\r', '\n')
        
        # Remove excessive blank lines (max 2 consecutive)
        text = re.sub(r'\n{3,}', '\n\n', text)
        
        # Remove null bytes and control characters (except newlines and tabs)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        
        # Strip trailing whitespace from each line
        lines = [line.rstrip() for line in text.split('\n')]
        text = '\n'.join(lines)
        
        return text.strip()
